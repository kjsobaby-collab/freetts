"""배포 매뉴얼·사용 설명서(.docx) 생성 스크립트.

사용: .venv\\Scripts\\python docs\\make_docs.py   (python-docx 필요: pip install -e ".[docx]")
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "src"))
from lecture_tts import __version__  # noqa: E402

DATE = "2026-10-21"
FONT = "Malgun Gothic"
MONO = "Consolas"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x59, 0x59, 0x59)
REPO = "https://github.com/kjsobaby-collab/freetts"  # main()에서 인자로 바꿀 수 있음


# ---------------------------------------------------------------- 공통 서식
def _font(run, name=FONT, size=None, bold=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade(el, fill):
    """문단(pPr) 또는 셀(tcPr)에 배경색."""
    pr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _field(paragraph, instr):
    run = paragraph.add_run()
    _font(run, size=9, color=GRAY)
    for tag, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


class Doc:
    def __init__(self, title: str, subtitle: str, audience: str):
        self.d = Document()
        self.title = title
        s = self.d.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin, s.bottom_margin = Cm(2.2), Cm(2.0)
        st = self.d.styles
        normal = st["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.25
        for lvl, size in ((1, 16), (2, 13), (3, 11)):
            h = st[f"Heading {lvl}"]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = NAVY
            rpr = h.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.insert(0, rf)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(a), FONT)
            h.paragraph_format.space_before = Pt(16 if lvl == 1 else 10)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        self._cover(subtitle, audience)

    # 표지 + 머리글/바닥글
    def _cover(self, subtitle, audience):
        d = self.d
        for _ in range(6):
            d.add_paragraph()
        p = d.add_paragraph()
        _font(p.add_run("Lecture TTS"), size=14, bold=True, color=GRAY)
        p = d.add_paragraph()
        _font(p.add_run(self.title), size=28, bold=True, color=NAVY)
        p.paragraph_format.space_after = Pt(2)
        p = d.add_paragraph()
        _font(p.add_run(subtitle), size=12, color=GRAY)
        bar = d.add_paragraph()
        b = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "18"), ("w:space", "1"), ("w:color", "1F3A5F")):
            bottom.set(qn(k), v)
        b.append(bottom)
        bar._p.get_or_add_pPr().append(b)
        for _ in range(3):
            d.add_paragraph()
        self.table(["항목", "내용"], [
            ["대상 프로그램", f"Lecture TTS {__version__} (Windows 10/11 64비트)"],
            ["문서 버전", "1.0"],
            ["작성일", DATE],
            ["대상 독자", audience],
            ["배포 주소", REPO],
        ], [4, 12], header=False)
        self.page_break()
        sec = d.sections[0]
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _font(hp.add_run(f"Lecture TTS · {self.title}"), size=8, color=GRAY)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(fp, "PAGE")
        sec.different_first_page_header_footer = True

    # 본문 요소
    def h1(self, t):
        self.d.add_heading(t, 1)

    def h2(self, t):
        self.d.add_heading(t, 2)

    def h3(self, t):
        self.d.add_heading(t, 3)

    def _inline(self, p, text, size=None, color=None):
        """**굵게**, `코드` 인라인 서식."""
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not part:
                continue
            if part.startswith("**"):
                _font(p.add_run(part[2:-2]), size=size, bold=True, color=color)
            elif part.startswith("`"):
                r = p.add_run(part[1:-1])
                _font(r, MONO, size=(size or 10) - 0.5, color=RGBColor(0xA3, 0x15, 0x15))
            else:
                _font(p.add_run(part), size=size, color=color)
        return p

    def p(self, text, size=None, color=None):
        return self._inline(self.d.add_paragraph(), text, size, color)

    def bullets(self, items, level=0):
        for it in items:
            p = self.d.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
            p.paragraph_format.space_after = Pt(2)
            self._inline(p, it)

    def steps(self, items):
        """번호 목록 (Word 자동 번호는 목록끼리 이어지므로 직접 번호를 붙인다)."""
        for i, it in enumerate(items, 1):
            p = self.d.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.8)
            pf.first_line_indent = Cm(-0.6)
            pf.space_after = Pt(3)
            _font(p.add_run(f"{i}.  "), bold=True, color=NAVY)
            self._inline(p, it)

    def code(self, lines):
        if isinstance(lines, str):
            lines = lines.strip("\n").splitlines()
        for i, line in enumerate(lines):
            p = self.d.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(4) if i == 0 else Pt(0)
            pf.space_after = Pt(4) if i == len(lines) - 1 else Pt(0)
            pf.line_spacing = 1.1
            pf.left_indent = Cm(0.3)
            _shade(p._p, "F2F2F2")
            _font(p.add_run(line or " "), MONO, size=9)

    def note(self, label, text, kind="info"):
        fill, color = {"info": ("EAF1FB", NAVY), "warn": ("FFF4E5", RGBColor(0xB2, 0x5E, 0x00)),
                       "danger": ("FDECEA", RGBColor(0xB0, 0x20, 0x20))}[kind]
        t = self.d.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0)
        c.width = Cm(16.6)
        _shade(c._tc, fill)
        p = c.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        _font(p.add_run(label + "  "), bold=True, color=color, size=9.5)
        texts = [text] if isinstance(text, str) else text
        self._inline(p, texts[0], size=9.5)
        for extra in texts[1:]:
            q = c.add_paragraph()
            q.paragraph_format.space_after = Pt(3)
            self._inline(q, extra, size=9.5)
        self.d.add_paragraph().paragraph_format.space_after = Pt(0)

    def table(self, headers, rows, widths_cm, header=True):
        t = self.d.add_table(rows=0, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for ci, w in enumerate(widths_cm):
            t.columns[ci].width = Cm(w)
        all_rows = ([headers] if header else []) + rows
        for ri, row in enumerate(all_rows):
            cells = t.add_row().cells
            for ci, val in enumerate(row):
                c = cells[ci]
                c.width = Cm(widths_cm[ci])
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.15
                if header and ri == 0:
                    _shade(c._tc, "1F3A5F")
                    _font(p.add_run(val), size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
                else:
                    if not header and ci == 0:
                        _shade(c._tc, "EEF2F7")
                    parts = str(val).split("\n")
                    self._inline(p, parts[0], size=9)
                    for extra in parts[1:]:
                        q = c.add_paragraph()
                        q.paragraph_format.space_after = Pt(1)
                        self._inline(q, extra, size=9)
        if header:
            tr = t.rows[0]._tr
            trpr = tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            th.set(qn("w:val"), "true")
            trpr.append(th)
        self.d.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    def image(self, path, width_cm=16.5, caption=None):
        self.d.add_picture(str(path), width=Cm(width_cm))
        self.d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            p = self.d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _font(p.add_run(caption), size=9, color=GRAY)

    def contents(self, items):
        self.h1("목차")
        for it in items:
            p = self.d.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            indent = it.startswith("  ")
            p.paragraph_format.left_indent = Cm(0.8 if indent else 0)
            _font(p.add_run(it.strip()), size=10 if indent else 10.5, bold=not indent,
                  color=None if indent else NAVY)
        self.page_break()

    def page_break(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        self.d.save(str(path))
        print("저장:", path)


# ---------------------------------------------------------------- 배포 매뉴얼
def deploy_manual(out: Path):
    v = __version__
    d = Doc("배포 매뉴얼", "다른 컴퓨터에 배포·설치·업데이트·제거하는 방법", "배포 담당자, 기관 IT 관리자")
    d.contents([
        "1. 개요", "2. 배포 구성", "3. 시스템 요구사항",
        "4. 방법 A — 설치 파일 만들기 (빌드 PC)",
        "5. 방법 A — 설치 파일 배포하기",
        "6. 방법 A — 사용자 PC에 설치하기",
        "7. 업데이트와 제거",
        "8. 방법 B — 소스 코드로 설치하기 (GitHub)",
        "9. 배포 전 점검 (릴리스 체크리스트)",
        "10. 라이선스 고지",
        "11. 문제 해결",
        "부록. 경로·파일 요약",
    ])

    d.h1("1. 개요")
    d.p("Lecture TTS는 강의 스크립트(.md/.txt/.docx/.srt)를 영어 내레이션 음성(mp3/wav)과 자막(SRT/VTT)으로 "
        "만드는 Windows 데스크톱 프로그램입니다. 이 문서는 개발 PC가 아닌 **다른 컴퓨터**에 프로그램을 "
        "배포하고 설치하는 절차를 설명합니다.")
    d.h2("1.1 배포 방법 선택")
    d.table(["구분", "방법 A — 설치 파일 (권장)", "방법 B — 소스 코드 설치"], [
        ["대상 사용자", "일반 사용자 (Python 지식 불필요)", "개발자, 시험 사용자"],
        ["사용자 준비물", "설치 파일 하나", "Python 3.12, ffmpeg, 명령줄 사용"],
        ["관리자 권한", "필요 없음 (사용자별 설치)", "필요 없음"],
        ["GPU 가속", "사용 안 함 (CPU 전용 빌드)", "NVIDIA GPU 사용 가능"],
        ["배포 형태", "GitHub Releases에 .exe 첨부", "GitHub 저장소에서 다운로드"],
        ["업데이트", "새 설치 파일을 덮어 설치", "git pull 또는 새로 다운로드"],
    ], [3.2, 6.7, 6.7])
    d.note("검증 상태", [
        f"버전 {__version__} 설치 파일은 {DATE}에 빌드해 개발 PC에서 설치 → 첫 실행 모델 다운로드 → 합성(mp3·자막) → GUI 실행 → "
        "업데이트 설치 → 제거까지 확인했습니다.",
        "깨끗한 Windows 10/11 PC, 한글 사용자 이름 계정, 오프라인 설치판은 아직 확인하지 않았습니다. 9장 체크리스트로 확인하세요.",
    ], "info")

    d.h1("2. 배포 구성")
    d.h2("2.1 설치 후 폴더 구성 (방법 A)")
    d.table(["경로", "내용"], [
        ["`%LOCALAPPDATA%\\Programs\\LectureTTS\\`", "앱 폴더 (설치 위치)"],
        ["  `LectureTTS.exe`", "실행 파일"],
        ["  `_internal\\`", "Python 런타임, PySide6(Qt), PyTorch(CPU), Kokoro, espeak-ng 라이브러리 등"],
        ["  `bin\\ffmpeg.exe`", "mp3 인코딩용 ffmpeg (LGPL 빌드)"],
        ["  `samples\\lecture_sample.md`", "예제 스크립트"],
        ["  `models\\`", "Kokoro 음성 모델 (오프라인 설치판에만 포함)"],
        ["  `THIRD_PARTY_LICENSES.txt`", "오픈소스 라이선스 고지"],
        ["`%LOCALAPPDATA%\\LectureTTS\\`", "사용자 데이터 폴더 — 업데이트해도 유지"],
        ["  `settings.yaml`, `lexicon.yaml`", "사용자 설정, 발음 사전"],
        ["  `models\\`", "첫 실행 때 내려받은 Kokoro 모델 (약 330MB)"],
        ["  `cache\\tts\\`", "문장 단위 합성 캐시"],
        ["  `logs\\`", "로그 파일"],
    ], [7, 9.6])
    d.p("앱 폴더와 사용자 데이터 폴더가 분리되어 있어, 새 버전을 설치해도 설정·발음 사전·캐시·모델은 그대로 남습니다.")
    d.h2("2.2 설치 파일 종류")
    d.table(["파일", "용도"], [
        [f"`LectureTTS-{v}-setup.exe`", "온라인 설치판. 첫 실행 때 모델(약 330MB)을 인터넷에서 내려받습니다."],
        [f"`LectureTTS-{v}-offline-setup.exe`", "오프라인 설치판. 모델이 포함되어 인터넷 없이 바로 쓸 수 있습니다 (학교·기관용)."],
        ["`*.exe.sha256`", "설치 파일의 SHA-256 해시. 내려받은 파일이 변조되지 않았는지 확인할 때 씁니다."],
    ], [6.5, 10.1])

    d.h1("3. 시스템 요구사항")
    d.h2("3.1 사용자 PC (설치 대상)")
    d.table(["항목", "요구사항"], [
        ["운영체제", "Windows 10 / 11 64비트"],
        ["메모리", "8GB 이상 권장"],
        ["디스크", "약 2GB 이상 여유 공간 (앱 약 1.0GB + 모델 약 330MB + 캐시)"],
        ["다운로드", f"설치 파일 약 252MB ({REPO}/releases)"],
        ["GPU", "필요 없음 (설치판은 CPU로 합성)"],
        ["인터넷", "첫 실행 때 모델 다운로드에만 필요 (오프라인 설치판은 불필요)"],
        ["권한", "일반 사용자 계정으로 설치 가능 (관리자 권한 불필요)"],
    ], [3.5, 13.1])
    d.h2("3.2 빌드 PC (설치 파일을 만드는 PC)")
    d.table(["항목", "요구사항"], [
        ["운영체제", "Windows 10 / 11 64비트"],
        ["Python", "**3.12** (python.org 설치판, `py -3.12`로 실행 가능해야 함). 3.13 이상은 kokoro가 지원하지 않습니다."],
        ["Inno Setup", "**6.x** — https://jrsoftware.org/isdl.php (기본 경로 `C:\\Program Files (x86)\\Inno Setup 6`)"],
        ["인터넷", "PyTorch(CPU), 패키지, ffmpeg, (오프라인판) 모델 다운로드"],
        ["디스크", "빌드 환경과 결과물용 10GB 이상 여유 공간 권장"],
        ["폴더 위치", "**짧은 경로**(예: `D:\\lecture-tts`). 경로가 길면 Windows 경로 길이 제한(260자) 때문에 설치 파일 단계가 실패합니다."],
        ["소스 코드", f"GitHub 저장소: {REPO}"],
    ], [3.5, 13.1])

    d.h1("4. 방법 A — 설치 파일 만들기 (빌드 PC)")
    d.h2("4.1 소스 코드 받기")
    d.code([f"git clone {REPO}.git", "cd freetts"])
    d.p("git이 없으면 GitHub 저장소 페이지에서 **Code → Download ZIP**으로 받아 압축을 풉니다.")
    d.h2("4.2 버전 번호 정하기")
    d.p("새 버전을 배포할 때마다 아래 두 파일의 버전을 같은 값으로 바꿉니다. 형식은 `MAJOR.MINOR.PATCH`(예: 1.2.0)입니다. "
        "이 번호는 설치 파일 이름과 앱의 **도움말 → 정보** 화면에 표시됩니다.")
    d.table(["파일", "수정할 줄"], [
        ["`pyproject.toml`", f'`version = "{v}"`'],
        ["`src\\lecture_tts\\__init__.py`", f'`__version__ = "{v}"`'],
    ], [6.5, 10.1])
    d.h2("4.3 빌드 실행")
    d.p("프로젝트 폴더에서 PowerShell을 열고 실행합니다.")
    d.code([
        "# 온라인 설치판",
        "powershell -ExecutionPolicy Bypass -File packaging\\build.ps1",
        "",
        "# 오프라인 설치판 (모델 포함)",
        "powershell -ExecutionPolicy Bypass -File packaging\\build.ps1 -Offline",
        "",
        "# 빌드 환경(.venv-build)이 이미 있으면 패키지 설치를 건너뜀",
        "powershell -ExecutionPolicy Bypass -File packaging\\build.ps1 -SkipVenv",
    ])
    d.p("`build.ps1`은 다음 순서로 진행합니다. 처음 빌드할 때는 패키지 다운로드 때문에 시간이 오래 걸립니다.")
    d.steps([
        "`.venv-build` 가상환경을 만들고 **CPU용 PyTorch**를 설치합니다. 어느 PC에서나 동작하고 용량도 줄어듭니다.",
        "Lecture TTS와 kokoro·PySide6·python-docx, PyInstaller, pip-licenses, spaCy 영어 모델(en_core_web_sm)을 설치합니다.",
        "`packaging\\bin\\ffmpeg.exe`가 없으면 `fetch_ffmpeg.ps1`로 **LGPL 빌드** ffmpeg(BtbN)를 내려받습니다.",
        "PyInstaller(onedir)로 `dist\\LectureTTS\\` 앱 폴더를 만듭니다.",
        "ffmpeg, `THIRD_PARTY_LICENSES.txt`, 예제 스크립트를 앱 폴더에 복사합니다.",
        "`-Offline`이면 Kokoro 모델을 `dist\\LectureTTS\\models\\`에 내려받습니다.",
        "앱 폴더 크기를 출력합니다. 이 값을 3.1절 디스크 요구사항에 적습니다.",
        "Inno Setup으로 `dist\\installer\\` 아래에 설치 파일과 `.sha256` 해시 파일을 만듭니다.",
    ])
    d.h2("4.4 결과물")
    d.table(["경로", "내용"], [
        ["`dist\\LectureTTS\\`", "설치하지 않고 바로 실행할 수 있는 앱 폴더"],
        [f"`dist\\installer\\LectureTTS-{v}-setup.exe`", "온라인 설치 파일"],
        [f"`dist\\installer\\LectureTTS-{v}-offline-setup.exe`", "오프라인 설치 파일 (`-Offline` 빌드 시)"],
        ["`dist\\installer\\*.exe.sha256`", "SHA-256 해시"],
    ], [8.5, 8.1])
    d.note("Inno Setup이 없을 때", [
        "설치 파일 단계만 건너뛰고 경고를 출력합니다. 이때는 `dist\\LectureTTS` 폴더를 zip으로 압축해 배포해도 됩니다 "
        "(휴대용 배포).",
        "사용자는 압축을 풀고 `LectureTTS.exe`를 실행하면 됩니다. 이 방식에는 시작 메뉴 바로가기와 제거 기능이 없습니다.",
    ])
    d.p("실측값(0.1.0, CPU용 PyTorch): 앱 폴더 약 1,000MB, 온라인 설치 파일 약 252MB. 처음 빌드는 패키지 다운로드를 포함해 약 20~30분 걸립니다.")
    d.h2("4.5 빌드 PC에서 먼저 확인")
    d.p("먼저 자동 점검을 실행합니다. 예제 스크립트의 첫 섹션을 실제 엔진으로 합성해 결과를 파일로 남깁니다 "
        "(모델이 없으면 먼저 내려받으므로 인터넷이 필요합니다).")
    d.code([
        "dist\\LectureTTS\\LectureTTS.exe --selftest D:\\selftest",
        "Get-Content D:\\selftest\\selftest.txt     # 첫 줄이 OK이면 정상",
    ])
    d.p("그다음 화면으로 확인합니다.")
    d.steps([
        "`dist\\LectureTTS\\LectureTTS.exe`를 실행합니다.",
        "**파일 → 예제 스크립트 열기** → 문장 하나를 선택하고 **선택 영역 듣기**를 누릅니다.",
        "**▶ 전체 생성** → **출력 폴더 열기**를 눌러 mp3와 srt가 만들어졌는지 확인합니다.",
    ])

    d.h1("5. 방법 A — 설치 파일 배포하기")
    d.h2("5.1 GitHub Releases에 올리기")
    d.steps([
        f"저장소 페이지({REPO})에서 **Releases → Draft a new release**를 누릅니다.",
        f"태그를 `v{v}` 형식으로 만듭니다. 버전 번호와 같아야 앱의 업데이트 확인이 정확히 동작합니다.",
        "릴리스 노트를 씁니다: 바뀐 점, 알려진 문제, 설치 방법(이 문서 6장 요약).",
        "`setup.exe`, `offline-setup.exe`, `.sha256` 파일을 첨부하고 **Publish release**를 누릅니다.",
        "릴리스 노트 본문에 각 설치 파일의 SHA-256 값을 함께 적습니다.",
    ])
    d.p("GitHub CLI를 쓰면 한 줄로 올릴 수 있습니다.")
    d.code(f'gh release create v{v} dist\\installer\\* --title "Lecture TTS {v}" --notes-file RELEASE_NOTES.md')
    d.note("용량 제한", "GitHub Releases는 파일 하나당 2GB까지 올릴 수 있습니다. 오프라인 설치판이 이보다 크면 "
           "기관 내부 파일 서버나 클라우드 드라이브로 배포하세요.", "warn")
    d.h2("5.2 사용자에게 안내할 내용")
    d.bullets([
        f"다운로드 주소: {REPO}/releases/latest",
        "인터넷이 막힌 PC는 **offline-setup.exe**를 받도록 안내",
        "SmartScreen 경고가 뜰 때 실행하는 방법 (6.2절)",
        "프로그램 사용법: 별도 문서 「Lecture TTS 사용 설명서」",
    ])
    d.h2("5.3 해시 확인 방법 (사용자)")
    d.code(f"Get-FileHash .\\LectureTTS-{v}-setup.exe -Algorithm SHA256")
    d.p("출력된 값이 릴리스 페이지의 SHA-256 값과 같으면 정상 파일입니다.")

    d.h1("6. 방법 A — 사용자 PC에 설치하기")
    d.h2("6.1 설치")
    d.steps([
        f"`LectureTTS-{v}-setup.exe`를 실행합니다.",
        "Windows 보안 경고가 뜨면 6.2절을 따릅니다.",
        "라이선스 안내(오픈소스 고지)를 확인하고 **동의**합니다.",
        "설치 위치는 기본값(`%LOCALAPPDATA%\\Programs\\LectureTTS`)을 그대로 둡니다. 관리자 권한이 필요 없습니다.",
        "필요하면 **바탕화면 바로가기 만들기**를 선택하고 **설치**를 누릅니다.",
        "마지막 화면에서 **Lecture TTS 실행**을 체크한 채 **마침**을 누릅니다.",
    ])
    d.p("설치하면 시작 메뉴에 **Lecture TTS** 폴더가 생기고, 그 안에 실행·예제 스크립트·제거 바로가기가 들어갑니다.")
    d.h2("6.2 Windows SmartScreen 경고")
    d.p("코드 서명을 하지 않은 설치 파일은 처음 실행할 때 \"Windows의 PC 보호\" 경고가 뜹니다.")
    d.steps([
        "경고 창에서 **추가 정보**를 누릅니다.",
        "게시자가 \"알 수 없는 게시자\"로 표시됩니다. 파일 이름이 맞는지 확인합니다.",
        "**실행**을 누릅니다.",
    ])
    d.note("참고", "백신이 PyInstaller로 만든 실행 파일을 의심 파일로 잘못 판정하는 경우가 있습니다. 해시가 릴리스 페이지와 "
           "같은지 확인한 뒤 예외로 등록하세요. 배포 규모가 커지면 코드 서명 인증서로 서명하는 것을 권장합니다.")
    d.h2("6.3 첫 실행 — 초기 설정 마법사")
    d.p("처음 실행하면 **초기 설정** 창이 열리고 필요한 구성 요소를 점검합니다.")
    d.table(["점검 항목", "정상 상태", "문제가 있을 때"], [
        ["ffmpeg", "앱에 포함된 `bin\\ffmpeg.exe`", "앱을 다시 설치합니다."],
        ["espeak-ng", "내장 라이브러리(espeakng-loader)", "설치판에 포함되어 있어 따로 설치할 필요가 없습니다."],
        ["Kokoro 모델", "준비됨", "**모델 다운로드**를 누릅니다 (약 330MB). 중간에 끊겨도 다시 누르면 이어받습니다."],
    ], [3.3, 5.5, 7.8])
    d.p("모든 항목이 ✓이면 **닫기**를 누릅니다. 마법사는 나중에 **도구 → 초기 설정 마법사…**에서 다시 열 수 있습니다.")
    d.h2("6.4 인터넷이 막힌 PC")
    d.p("다음 두 방법 중 하나를 씁니다.")
    d.bullets([
        "**오프라인 설치판** 사용: 모델이 앱 폴더에 들어 있어 별도 작업이 필요 없습니다.",
        "**모델 폴더 가져오기**: 인터넷이 되는 PC에서 모델 폴더 `models--hexgrad--Kokoro-82M`을 USB 등으로 옮긴 뒤, "
        "초기 설정 마법사의 **모델 폴더 가져오기…**에서 그 폴더를 선택합니다.",
    ])
    d.table(["모델 폴더가 있는 위치 (복사 원본)", "경로"], [
        ["Lecture TTS 설치판을 쓰는 PC", "`%LOCALAPPDATA%\\LectureTTS\\models\\hub\\models--hexgrad--Kokoro-82M`"],
        ["Python 개발 환경(방법 B)을 쓰는 PC", "`%USERPROFILE%\\.cache\\huggingface\\hub\\models--hexgrad--Kokoro-82M`"],
    ], [6, 10.6])
    d.h2("6.5 설치 확인")
    d.steps([
        "**파일 → 예제 스크립트 열기**를 누릅니다.",
        "편집기에서 문장 하나를 드래그해 선택하고 오른쪽 **▶ 선택 영역 듣기**를 누릅니다. 소리가 나면 정상입니다.",
        "**▶ 전체 생성**을 누르고, 끝나면 **출력 폴더 열기**로 결과 파일을 확인합니다.",
    ])
    d.h2("6.6 여러 PC에 무인 설치 (IT 관리자)")
    d.p("Inno Setup 표준 옵션을 쓸 수 있습니다. 배포 도구나 로그온 스크립트에서 사용자 권한으로 실행합니다.")
    d.code(f'LectureTTS-{v}-offline-setup.exe /VERYSILENT /SUPPRESSMSGBOXES /NORESTART /TASKS="desktopicon"')

    d.h1("7. 업데이트와 제거")
    d.h2("7.1 업데이트")
    d.bullets([
        "새 버전 설치 파일을 그대로 실행해 **덮어 설치**합니다. 기존 버전을 먼저 지울 필요가 없습니다.",
        "설정·발음 사전·캐시·모델(`%LOCALAPPDATA%\\LectureTTS`)은 유지됩니다.",
        "예전 버전에서 만든 프로젝트 파일(.lttproj)은 열 때 자동으로 변환됩니다.",
    ])
    d.h2("7.2 새 버전 알림 설정")
    d.p("앱은 시작할 때 배포 페이지의 최신 버전을 확인하고 **알림만** 표시합니다. 자동으로 설치하지는 않습니다. "
        "확인 주소는 **도구 → 환경설정… → 업데이트**에서 지정합니다. GitHub Releases를 쓰면 아래 주소를 넣습니다.")
    d.code(REPO.replace("https://github.com/", "https://api.github.com/repos/") + "/releases/latest")
    d.p("이 주소는 기본값으로 들어 있습니다(`src\\lecture_tts\\__init__.py`의 `UPDATE_URL`). 저장소를 옮기면 이 값을 바꾼 뒤 빌드합니다.")
    d.h2("7.3 제거")
    d.steps([
        "**설정 → 앱 → 설치된 앱**에서 **Lecture TTS**를 제거하거나, 시작 메뉴의 **Lecture TTS 제거**를 실행합니다.",
        "앱 폴더는 자동으로 삭제됩니다.",
        "\"설정·발음 사전·캐시·모델도 삭제할까요?\"라는 질문이 나옵니다. 완전히 지우려면 **예**, 다시 설치할 예정이면 **아니요**를 누릅니다.",
    ])
    d.p("무인 제거 시에는 질문 없이 사용자 데이터를 남겨 둡니다.")
    d.code('"%LOCALAPPDATA%\\Programs\\LectureTTS\\unins000.exe" /VERYSILENT /SUPPRESSMSGBOXES')

    d.h1("8. 방법 B — 소스 코드로 설치하기 (GitHub)")
    d.p("Python 개발 환경에서 직접 실행하는 방법입니다. 설치 파일이 아직 없거나 NVIDIA GPU로 빠르게 합성하고 싶을 때 씁니다.")
    d.h2("8.1 준비")
    d.table(["항목", "설치 방법"], [
        ["Python 3.12", "https://www.python.org/downloads/ 에서 3.12.x 설치. 설치 후 `py -3.12 --version`으로 확인"],
        ["git (선택)", "https://git-scm.com/ — 없으면 ZIP으로 내려받습니다."],
        ["ffmpeg", "8.2절 4단계의 스크립트로 프로젝트 안에 내려받습니다. PATH에 ffmpeg가 이미 있으면 생략할 수 있습니다."],
    ], [3.3, 13.3])
    d.h2("8.2 설치")
    d.p("PowerShell에서 실행합니다.")
    d.code([
        "# 1. 소스 받기",
        f"git clone {REPO}.git",
        "cd freetts",
        "",
        "# 2. 가상환경",
        "py -3.12 -m venv .venv",
        ".venv\\Scripts\\python -m pip install --upgrade pip",
        "",
        "# 3. PyTorch — 둘 중 하나",
        ".venv\\Scripts\\pip install torch --index-url https://download.pytorch.org/whl/cpu   # CPU",
        "#   NVIDIA GPU: https://pytorch.org 의 'Get Started'에서 CUDA용 설치 명령을 확인해 실행",
        "",
        "# 4. 프로그램과 의존 패키지, ffmpeg",
        ".venv\\Scripts\\pip install -e \".[kokoro,gui,docx]\"",
        ".venv\\Scripts\\pip install https://github.com/explosion/spacy-models/releases/download/"
        "en_core_web_sm-3.8.0/en_core_web_sm-3.8.0-py3-none-any.whl",
        "powershell -ExecutionPolicy Bypass -File packaging\\fetch_ffmpeg.ps1",
        "",
        "# 5. 모델 다운로드 (최초 1회, 약 330MB)",
        ".venv\\Scripts\\lecture-tts setup kokoro",
    ])
    d.h2("8.3 실행")
    d.code([
        ".venv\\Scripts\\lecture-tts-gui                      # GUI",
        ".venv\\Scripts\\lecture-tts build my_lecture.md      # 명령줄",
    ])
    d.p("업데이트는 프로젝트 폴더에서 `git pull`을 실행한 뒤 4단계 `pip install -e` 명령을 다시 실행합니다.")
    d.note("GPU 확인", "`.venv\\Scripts\\python -c \"import torch; print(torch.cuda.is_available())\"`가 True이면 GPU로 합성합니다. "
           "RTX 3080 기준으로 문장 16개(슬라이드 3장)를 약 36초에 생성했습니다.")

    d.h1("9. 배포 전 점검 (릴리스 체크리스트)")
    d.p("설치 파일을 공개하기 전에 아래 항목을 모두 확인합니다. 개발 도구가 깔린 PC에서는 숨은 문제가 드러나지 않으므로 "
        "**깨끗한 PC 또는 가상머신**에서 시험합니다.")
    d.table(["☐", "점검 항목", "확인 방법"], [
        ["☐", "깨끗한 Windows 10 PC 설치·실행", "설치 → 첫 실행 → 모델 다운로드 → 생성 → 재생"],
        ["☐", "깨끗한 Windows 11 PC 설치·실행", "위와 같음"],
        ["☐", "관리자 권한 없는 계정", "일반 사용자 계정으로 설치·생성"],
        ["☐", "한글 사용자 이름 계정", "사용자 폴더 경로에 한글이 있는 계정으로 설치·생성"],
        ["☐", "한글·공백이 있는 경로", "\"D:\\강의 자료\\3주차.md\" 같은 파일 열기·생성"],
        ["☐", "오프라인 설치판", "인터넷을 끊은 상태에서 설치·생성"],
        ["☐", "업데이트 설치", "이전 버전 위에 설치한 뒤 설정·발음 사전 유지 확인"],
        ["☐", "제거", "앱 폴더 삭제, 사용자 데이터 삭제 여부 질문 확인"],
        ["☐", "라이선스 파일", "앱 폴더의 `THIRD_PARTY_LICENSES.txt` 내용 확인"],
        ["☐", "버전 번호", "설치 파일 이름, 도움말 → 정보, 릴리스 태그가 모두 같은지"],
        ["☐", "해시값", "릴리스 노트의 SHA-256과 실제 파일 비교"],
        ["☐", "릴리스 노트", "바뀐 점과 알려진 문제 작성"],
        ["☐", "자동 테스트", "`.venv\\Scripts\\pytest` 전체 통과"],
    ], [1, 5.6, 10])

    d.h1("10. 라이선스 고지")
    d.p("배포판에는 여러 오픈소스 구성 요소가 포함됩니다. 고지는 앱의 **도움말 → 정보 → 오픈소스 라이선스** 탭과 "
        "설치 폴더의 `THIRD_PARTY_LICENSES.txt`에 모여 있습니다.")
    d.table(["구성 요소", "라이선스", "배포 시 확인 사항"], [
        ["Kokoro 모델·코드", "Apache 2.0", "고지 포함"],
        ["PySide6 / Qt", "LGPL v3", "onedir 빌드로 Qt DLL을 동적 링크 상태로 유지(사용자가 교체 가능), 라이선스 원문 포함"],
        ["ffmpeg", "LGPL (BtbN lgpl 빌드)", "`bin\\FFMPEG_SOURCE.txt`, `FFMPEG_LICENSE.txt`에 출처·라이선스 기재"],
        ["espeak-ng", "GPL v3", "espeakng-loader 패키지에 포함되어 함께 배포됨. **GPL 조건이 배포 방식에 미치는 영향을 배포 전 검토**"],
        ["PyTorch 등 Python 패키지", "각 패키지 라이선스", "pip-licenses로 목록 자동 생성"],
    ], [3.8, 3.4, 9.4])
    d.note("배포 전 결정할 사항", [
        "① espeak-ng(GPL) 동봉 여부 ② 코드 서명 인증서 도입 여부 ③ 앱 아이콘(`packaging\\icon.ico`) ④ 기본 업데이트 확인 주소.",
        "라이선스 해석이 애매한 항목은 소속 기관의 담당 부서나 전문가에게 확인하세요.",
    ], "warn")

    d.h1("11. 문제 해결")
    d.table(["증상", "원인", "해결"], [
        ["설치 파일 실행 시 \"Windows의 PC 보호\"", "코드 서명이 없음", "추가 정보 → 실행 (6.2절)"],
        ["백신이 설치 파일·exe를 차단", "PyInstaller 실행 파일 오탐", "해시 확인 후 예외 등록"],
        ["모델 다운로드 실패", "인터넷·방화벽 차단", "다시 누르면 이어받기. 막힌 환경은 오프라인판 또는 모델 폴더 가져오기 (6.4절)"],
        ["상태 줄에 \"초기 설정 필요\"", "모델 없음", "도구 → 초기 설정 마법사 → 모델 다운로드"],
        ["mp3가 만들어지지 않음", "ffmpeg 없음", "방법 A: 재설치. 방법 B: `packaging\\fetch_ffmpeg.ps1` 실행"],
        ["합성이 매우 느림", "CPU 합성", "정상입니다. 방법 B에서 GPU용 PyTorch를 쓰면 빨라집니다."],
        ["빌드 중 `py -3.12`를 찾을 수 없음", "Python 3.12 미설치", "python.org에서 3.12 설치 (py launcher 포함)"],
        ["빌드 후 설치 파일이 없음", "Inno Setup 미설치", "Inno Setup 6 설치 후 `build.ps1 -SkipVenv`로 다시 실행"],
        ["Inno Setup 단계에서 \"지정된 경로를 찾을 수 없습니다\"", "경로 길이 제한(260자) 초과", "프로젝트를 짧은 경로(예: `D:\\lecture-tts`)로 옮겨 `build.ps1 -SkipVenv` 실행"],
        ["업데이트 후 \"DLL load failed while importing QtCore\"", "이전 버전 파일이 남음", "설치 파일이 `_internal` 폴더를 지우고 다시 설치합니다. 계속되면 제거 후 다시 설치"],
        ["설치 후 합성이 되는지 확인하고 싶음", "—", "`LectureTTS.exe --selftest D:\\selftest` 실행 후 `selftest.txt` 첫 줄이 OK인지 확인"],
        ["Git Bash에서 한글이 깨짐", "콘솔 인코딩(cp949)", "PowerShell이나 명령 프롬프트에서 실행"],
        ["프로그램이 비정상 종료", "—", "도움말 → 진단 정보 내보내기로 만든 zip을 개발자에게 전달 (스크립트 원문·API 키는 포함되지 않음)"],
    ], [4.6, 3.6, 8.4])

    d.h1("부록. 경로·파일 요약")
    d.table(["항목", "경로"], [
        ["앱 설치 폴더", "`%LOCALAPPDATA%\\Programs\\LectureTTS`"],
        ["사용자 데이터", "`%LOCALAPPDATA%\\LectureTTS`"],
        ["로그", "`%LOCALAPPDATA%\\LectureTTS\\logs`"],
        ["사용자 데이터 위치 변경", "환경 변수 `LECTURE_TTS_HOME`"],
        ["모델 위치 변경", "환경 변수 `LECTURE_TTS_MODELS`"],
        ["빌드 스크립트", "`packaging\\build.ps1`"],
        ["PyInstaller 설정", "`packaging\\lecture_tts_gui.spec`"],
        ["설치 파일 스크립트", "`packaging\\installer.iss`"],
        ["ffmpeg 다운로드", "`packaging\\fetch_ffmpeg.ps1`"],
        ["기본 설정", "`src\\lecture_tts\\data\\default.yaml`"],
        ["설계 문서", "`lecture tts design.md` (8.10절 외부 배포 요구사항)"],
    ], [5, 11.6])
    d.save(out)


# ---------------------------------------------------------------- 사용 설명서
def user_manual(out: Path, screenshot: Path | None):
    d = Doc("사용 설명서", "강의 스크립트로 영어 내레이션과 자막 만들기", "강의 음성을 만드는 사용자")
    d.contents([
        "1. 소개", "2. 시작하기", "3. 화면 구성", "4. 빠른 시작 — 처음 음성 만들기",
        "5. 스크립트 작성 규칙", "6. 음성 생성", "7. 듣기와 재생", "8. 설정 패널",
        "9. 도구 메뉴", "10. 결과 파일", "11. 프로젝트 파일", "12. 명령줄(CLI) 사용법",
        "13. 자주 묻는 질문", "부록. 음성 목록",
    ])

    d.h1("1. 소개")
    d.p("Lecture TTS는 강의 스크립트를 자연스러운 **영어 내레이션 음성**과 **자막**으로 바꿔 주는 프로그램입니다.")
    d.bullets([
        "입력: Markdown(.md), 텍스트(.txt), Word(.docx), 자막(.srt/.vtt)",
        "출력: 섹션(슬라이드)별 음성 + 전체 병합 음성(mp3/wav), SRT/VTT 자막",
        "기본 엔진 **Kokoro**는 PC에서 직접 돌아가므로 인터넷 연결과 사용료가 필요 없습니다.",
        "문장 단위로 결과를 저장(캐시)해 두므로, 스크립트에서 한 문장만 고치면 그 문장만 다시 만듭니다.",
        "숫자·연도·약어·기호를 읽기 좋은 영어로 자동 변환합니다 (예: 2024 → twenty twenty-four).",
    ])

    d.h1("2. 시작하기")
    d.h2("2.1 실행")
    d.table(["설치 방식", "실행 방법"], [
        ["설치 파일로 설치", "시작 메뉴 → **Lecture TTS** (또는 바탕화면 바로가기)"],
        ["소스 코드로 설치", "프로젝트 폴더에서 `.venv\\Scripts\\lecture-tts-gui`"],
    ], [4, 12.6])
    d.h2("2.2 처음 실행할 때")
    d.p("**초기 설정** 창이 열리면 **모델 다운로드**를 눌러 음성 모델(약 330MB)을 받습니다. 한 번만 하면 됩니다. "
        "인터넷이 없는 PC는 관리자에게 오프라인 설치판을 요청하세요.")
    d.p("프로그램은 마지막으로 연 파일을 자동으로 다시 엽니다. 처음에는 예제 스크립트가 열립니다.")

    d.h1("3. 화면 구성")
    if screenshot and screenshot.exists():
        d.image(screenshot, 16.5, "그림 1. Lecture TTS 주 화면")
    d.table(["위치", "이름", "기능"], [
        ["맨 위", "메뉴", "파일 · 생성 · 도구 · 도움말"],
        ["메뉴 아래", "도구 막대", "열기 · 저장 · ▶ 전체 생성 · 변경분만 생성 · ■ 중지 · 출력 폴더 열기"],
        ["왼쪽", "섹션 목록", "슬라이드(섹션) 목록과 상태 표시. 클릭하면 편집기가 그 위치로 이동"],
        ["가운데", "편집기", "**원문 편집** · **정규화 미리보기**(실제로 읽을 문장) · **청크 보기**(합성 단위) 탭"],
        ["오른쪽", "설정 패널", "엔진 · 억양 · 음성 · 속도 · 포맷 · 출력 폴더, 선택 영역 듣기 · 음성 비교 · 발음 사전"],
        ["아래", "플레이어", "재생/정지, 재생 대상(전체 또는 섹션) 선택, 재생 위치, 음량"],
        ["맨 아래", "상태 줄 · 로그", "진행 상황과 남은 시간. **▸ 로그**를 누르면 자세한 기록이 펼쳐짐"],
        ["맨 아래 끝", "요약", "섹션 수 · 청크 수 · 글자 수 · 예상 길이 · 경고 수"],
    ], [2.6, 2.8, 11.2])
    d.h2("3.1 섹션 상태 표시")
    d.table(["표시", "의미"], [
        ["● 완료", "음성 파일이 최신 상태"],
        ["◐ 생성 중", "지금 합성하는 중"],
        ["○ 변경됨 / 미생성", "아직 만들지 않았거나, 만든 뒤 스크립트가 바뀜"],
        ["✕ 오류", "일부 문장 합성 실패"],
    ], [4, 12.6])

    d.h1("4. 빠른 시작 — 처음 음성 만들기")
    d.steps([
        "**파일 → 열기**(또는 도구 막대 **열기**)로 강의 스크립트를 엽니다.",
        "왼쪽 **섹션 목록**에 슬라이드가 제대로 나뉘었는지 확인합니다. 안 나뉘면 5장의 섹션 표기 규칙을 따릅니다.",
        "**정규화 미리보기** 탭에서 숫자·약어가 어떻게 읽히는지 확인합니다.",
        "오른쪽 **설정**에서 억양(미국/영국), 음성, 속도를 고릅니다.",
        "편집기에서 문장 몇 개를 드래그해 선택하고 **▶ 선택 영역 듣기**로 미리 들어 봅니다. 파일은 저장되지 않습니다.",
        "**▶ 전체 생성**을 누릅니다. 상태 줄에 진행률과 남은 시간이 표시됩니다.",
        "끝나면 아래 플레이어의 **▶**로 들어 보고, **출력 폴더 열기**로 결과 파일을 확인합니다.",
        "스크립트를 고친 뒤에는 **변경분만 생성**을 누르면 바뀐 부분만 다시 만듭니다.",
    ])
    d.note("자동 저장", "편집한 내용은 2초 뒤 자동 저장되고, 생성을 시작하기 직전에도 저장됩니다.")

    d.h1("5. 스크립트 작성 규칙")
    d.p("일반 텍스트에 몇 가지 표기만 더하면 됩니다. 앱 안에서는 **도움말 → 스크립트 문법**에서 같은 내용을 볼 수 있습니다.")
    d.table(["문법", "의미", "예시"], [
        ["`# 제목`", "강의 제목 (읽지 않음)", "`# Lecture 3`"],
        ["`## Slide N: 제목`\n`[SLIDE N]`\n`Slide N. 제목`\n`## 제목`", "섹션 구분 → 섹션별 음성 파일", "`## Slide 2: Intro`\n`Slide 2. Intro`"],
        ["`[pause 1s]` `[pause 500ms]` `[pause]`", "무음 삽입 (길이를 생략하면 1초)", "`[pause 2s]`"],
        ["`{표기|읽기}`", "발음 지정 — 자막에는 표기, 음성은 읽기대로", "`{SQL|sequel}`"],
        ["`*단어*`", "강조 표시 (기호만 지우고 평소대로 읽음)", "`*important*`"],
        ["`[speaker: 이름]`", "이후 문단의 화자 전환 (student는 남성 음성)", "`[speaker: student]`"],
        ["`[rate 0.9]`", "이후 문단의 속도 변경", "`[rate 0.85]`"],
        ["`<!-- 내용 -->`", "주석 (읽지 않음)", "`<!-- 수정 필요 -->`"],
        ["빈 줄", "문단 구분 → 문단 사이 쉼", ""],
    ], [5.2, 6.4, 5])
    d.bullets([
        "첫 섹션 표기 앞에 있는 줄(머리말)은 읽지 않습니다. 그중 첫 줄을 강의 제목으로 씁니다.",
        "잘못 쓴 지시어는 편집기에 물결 밑줄로 표시되고, 그 지시어는 무시됩니다.",
        "숫자·연도·%·통화·약어(e.g., i.e.)·기호(≤, ≈, x^2)는 자동으로 읽기 좋은 영어로 바뀝니다.",
    ])
    d.h2("5.1 예시")
    d.code("""
# Language and Culture - Week 3

## Slide 1: Introduction
Hello everyone, and welcome back. Today we look at K-pop and heritage.
[pause]

## Slide 2: Key Numbers
Sales grew by 45% in 2024. That is almost *double* the year before.
We will use {SQL|sequel} later in the course.

[speaker: student]
Can you explain that again?
""")

    d.h1("6. 음성 생성")
    d.table(["기능", "방법", "설명"], [
        ["전체 생성", "도구 막대 **▶ 전체 생성**", "모든 섹션을 생성합니다. 이미 만든 문장은 캐시를 다시 씁니다."],
        ["변경분만 생성", "**변경분만 생성**", "수정됐거나 아직 만들지 않은 섹션만 합성합니다."],
        ["한 섹션만 생성", "섹션 목록에서 오른쪽 클릭 → **이 섹션만 생성**", "특정 슬라이드만 다시 만들 때"],
        ["중지", "**■ 중지**", "지금 합성 중인 문장이 끝나면 멈춥니다. 완료된 문장은 캐시에 남습니다."],
    ], [3.2, 5.6, 7.8])
    d.bullets([
        "일부 문장이 실패하면 안내 창이 뜹니다. **변경분만 생성**을 다시 누르면 실패한 문장만 다시 시도합니다.",
        "전체 병합 파일(`이름_full.mp3`)은 모든 섹션이 만들어진 뒤에 생깁니다.",
        "생성 중에도 편집기에서 다른 부분을 고칠 수 있습니다.",
    ])

    d.h1("7. 듣기와 재생")
    d.table(["기능", "방법"], [
        ["선택 영역 듣기", "편집기에서 문장을 선택 → 오른쪽 **▶ 선택 영역 듣기** 또는 오른쪽 클릭 → **▶ 선택 영역 듣기**"],
        ["전체·섹션 재생", "아래 플레이어에서 대상(전체 / 섹션 N)을 고르고 **▶**"],
        ["섹션 재생", "섹션 목록에서 오른쪽 클릭 → **이 섹션 재생**"],
        ["특정 문장부터 재생", "편집기에서 **Ctrl + 더블클릭** (또는 오른쪽 클릭 → 이 문장부터 재생)"],
        ["읽는 위치 표시", "재생 중에는 지금 읽는 문장이 편집기에서 **노란색**으로 표시됩니다."],
    ], [4, 12.6])

    d.h1("8. 설정 패널")
    d.table(["항목", "설명"], [
        ["엔진", "**Kokoro (로컬, 기본)** — 고품질, 무료\n**Piper (로컬, 경량)** — 빠른 초안용 (별도 설치 필요)\n"
                 "**Azure Speech F0 (클라우드)** — 본인 Azure 키 필요\n**Fake (테스트용)** — 소리 없이 동작 확인"],
        ["억양", "미국 / 영국. 고른 억양의 음성만 목록에 나옵니다."],
        ["음성", "기본 `af_heart`(미국 여성). 부록의 음성 목록 참고"],
        ["속도", "0.50 ~ 1.50 (1.00이 보통 속도)"],
        ["포맷", "mp3(기본, 용량 작음) / wav(무손실)"],
        ["출력", "결과를 저장할 폴더. 비워 두면 **스크립트가 있는 폴더\\output**"],
    ], [3, 13.6])

    d.h1("9. 도구 메뉴")
    d.h2("9.1 음성 비교")
    d.p("**도구 → 음성 비교…**: 여러 음성과 속도로 같은 문장을 만들어 표에서 바로 비교합니다. "
        "마음에 드는 행에서 **기본 음성으로 지정**을 누르면 설정에 반영됩니다. "
        "전문 용어·숫자·긴 문장이 섞인 30초 분량 문단으로 비교하는 것이 좋습니다.")
    d.h2("9.2 발음 사전")
    d.p("**도구 → 발음 사전…**: 계속 틀리게 읽는 단어를 표기/읽기 표로 등록합니다. 행마다 **▶ 듣기**로 확인할 수 있습니다. "
        "편집기에서 단어를 오른쪽 클릭 → **발음 사전에 추가**로도 등록할 수 있습니다. "
        "사전은 모든 스크립트에 적용되고, 대소문자를 구분하며 단어 단위로 바꿉니다.")
    d.table(["표기", "읽기", "효과"], [
        ["`Gyeongbokgung`", "`gyung bok goong`", "한국어 고유명사의 발음 보정"],
        ["`Seongdeok`", "`sung duck`", "위와 같음"],
    ], [4, 4, 8.6])
    d.p("한 스크립트에서만 바꾸려면 사전 대신 `{표기|읽기}` 문법을 씁니다.", size=9.5, color=GRAY)
    d.h2("9.3 환경설정")
    d.p("**도구 → 환경설정…**에서 바꾼 값은 다음 실행에도 유지됩니다.")
    d.table(["탭/항목", "설명"], [
        ["오디오", "문단 사이 쉼, 섹션 앞뒤 여백, [pause] 기본 길이, 음량 목표, mp3 비트레이트"],
        ["텍스트·자막", "대문자 약어를 한 글자씩 읽기(API → A P I), 섹션 제목도 읽기, 자막 형식, 자막 한 줄 최대 글자"],
        ["캐시", "**캐시 비우기** — 다음 생성 때 모든 문장을 다시 합성"],
        ["Azure Speech", "본인의 키와 지역(예: koreacentral). 키는 Windows 자격 증명 관리자에 저장됩니다."],
        ["업데이트", "시작할 때 새 버전 확인(알림만 표시), 확인 주소"],
    ], [3.4, 13.2])
    d.h2("9.4 초기 설정 마법사")
    d.p("**도구 → 초기 설정 마법사…**: ffmpeg, espeak-ng, 음성 모델 상태를 다시 점검하고 모델을 내려받거나 가져옵니다.")
    d.h2("9.5 도움말 메뉴")
    d.bullets([
        "**스크립트 문법** — 5장의 작성 규칙 요약",
        "**진단 정보 내보내기…** — 오류를 알릴 때 쓰는 zip 파일(로그·환경 정보). 스크립트 원문과 API 키는 들어가지 않습니다.",
        "**업데이트 확인** — 새 버전이 있으면 배포 페이지를 엽니다.",
        "**정보** — 버전, 오픈소스 라이선스, 사용자 데이터 폴더 위치",
    ])

    d.h1("10. 결과 파일")
    d.p("기본 저장 위치는 `스크립트 폴더\\output\\스크립트 이름\\`입니다.")
    d.code("""
output\\week3\\
├─ sections\\
│   ├─ 01_introduction.mp3      섹션별 음성
│   ├─ 01_introduction.srt      섹션별 자막
│   ├─ 02_key-numbers.mp3
│   └─ ...
├─ week3_full.mp3               전체 병합 음성
├─ week3_full.srt / .vtt        전체 자막
├─ normalized_script.txt        실제로 읽은 문장 (숫자·약어 변환 결과)
└─ manifest.json                생성 정보 (엔진, 음성, 속도, 섹션별 길이)
""")
    d.bullets([
        "섹션 파일 이름은 `번호_제목` 형식입니다. 제목에 영문이 없으면 번호만 씁니다.",
        "자막 파일은 동영상 편집기나 LMS에 음성과 함께 올릴 수 있습니다.",
    ])

    d.h1("11. 프로젝트 파일")
    d.p("**파일 → 프로젝트 저장**은 스크립트 위치와 이 강의에만 쓸 설정(음성, 속도, 출력 폴더 등)을 `.lttproj` 파일로 저장합니다. "
        "**파일 → 프로젝트 열기…**로 열면 그 설정이 그대로 적용됩니다. 강의마다 음성이나 속도를 다르게 쓸 때 편리합니다.")
    d.p("파일 메뉴의 다른 기능: **새 스크립트**, **최근 파일**, **다른 이름으로 저장…**, **예제 스크립트 열기**. "
        ".docx/.srt 파일은 텍스트로 가져오며, 수정하면 .md로 저장됩니다.")

    d.h1("12. 명령줄(CLI) 사용법")
    d.p("GUI 없이 PowerShell에서 쓸 수 있습니다. 여러 강의를 한꺼번에 만들 때 편리합니다. 소스 코드 설치(방법 B) 환경에서 "
        "`.venv\\Scripts\\activate`를 먼저 실행하세요.")
    d.table(["명령", "설명"], [
        ["`lecture-tts build week3.md`", "전체 생성 (mp3 + 자막)"],
        ["`lecture-tts build week3.md --sections 2,5-7`", "일부 섹션만"],
        ["`lecture-tts build week3.md -v am_michael -r 0.95 -f wav`", "음성·속도·형식 지정"],
        ["`lecture-tts build week3.md -o D:\\out`", "출력 폴더 지정"],
        ["`lecture-tts build week3.md --no-cache`", "캐시를 무시하고 모두 다시 합성"],
        ["`lecture-tts preview week3.md --show chunks`", "합성하지 않고 읽을 문장만 확인"],
        ["`lecture-tts batch .\\scripts --glob \"*.md\"`", "폴더 안의 스크립트를 모두 생성"],
        ["`lecture-tts audition --voices af_heart,bf_emma`", "음성 비교 샘플 만들기"],
        ["`lecture-tts voices`", "쓸 수 있는 음성 목록"],
        ["`lecture-tts setup kokoro`", "음성 모델 미리 받기"],
        ["`lecture-tts cache info` / `cache clear`", "캐시 크기 확인 / 비우기"],
        ["`lecture-tts usage`", "Azure 사용량 확인"],
        ["`lecture-tts gui`", "GUI 실행"],
    ], [8.2, 8.4])
    d.p("각 명령 뒤에 `--help`를 붙이면 전체 옵션을 볼 수 있습니다. Git Bash에서는 한글이 깨지므로 PowerShell이나 명령 프롬프트를 쓰세요.",
        size=9.5, color=GRAY)

    d.h1("13. 자주 묻는 질문")
    faq = [
        ("단어를 틀리게 읽어요.", "`{표기|읽기}`로 그 자리만 고치거나, 오른쪽 클릭 → **발음 사전에 추가**로 모든 스크립트에 적용합니다."),
        ("숫자를 어떻게 읽는지 미리 알고 싶어요.", "**정규화 미리보기** 탭에서 실제로 읽을 문장을 볼 수 있습니다."),
        ("한 문장만 고쳤는데 전부 다시 만들어야 하나요?", "아닙니다. **변경분만 생성**을 누르면 고친 문장만 다시 합성합니다."),
        ("전체 병합 파일(_full.mp3)이 안 생겨요.", "모든 섹션이 만들어져야 생깁니다. **▶ 전체 생성** 또는 **변경분만 생성**을 누르세요."),
        ("생성 도중 중지하면 처음부터 다시 해야 하나요?", "아닙니다. 완료된 문장은 캐시에 남아 다음 생성 때 그대로 씁니다."),
        ("상태 줄에 \"초기 설정 필요\" 또는 \"엔진 준비 안 됨\"이 떠요.", "**도구 → 초기 설정 마법사…**에서 모델을 내려받으세요."),
        ("머리말이 읽히지 않아요.", "첫 섹션 표기 앞의 줄은 일부러 읽지 않습니다. 읽으려면 그 부분을 섹션 안으로 옮기세요."),
        ("섹션 제목도 읽게 하고 싶어요.", "**환경설정 → 텍스트·자막 → 섹션 제목도 읽기**를 켭니다."),
        ("쉼이 너무 길거나 짧아요.", "`[pause 0.5s]`처럼 길이를 지정하거나, 환경설정에서 문단 사이 쉼과 [pause] 기본 길이를 바꿉니다."),
        ("캐시가 디스크를 많이 차지해요.", "**환경설정 → 캐시 비우기** 또는 `lecture-tts cache clear`. 다음 생성은 처음부터 합성합니다."),
        ("오류가 나요.", "**도움말 → 진단 정보 내보내기…**로 만든 zip을 담당자에게 보내 주세요. 로그는 `%LOCALAPPDATA%\\LectureTTS\\logs`에 있습니다."),
    ]
    for q, a in faq:
        p = d.d.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        _font(p.add_run("Q. " + q), bold=True, color=NAVY)
        ap = d.p("A. " + a)
        ap.paragraph_format.left_indent = Cm(0.5)
        ap.paragraph_format.space_after = Pt(6)

    d.h1("부록. 음성 목록 (Kokoro)")
    d.p("음성 이름의 앞 두 글자는 억양과 성별을 뜻합니다: **a**=미국, **b**=영국 / **f**=여성, **m**=남성.")
    d.table(["구분", "음성"], [
        ["미국 여성 (af_)", "af_heart(기본), af_alloy, af_aoede, af_bella, af_jessica, af_kore, af_nicole, af_nova, af_river, af_sarah, af_sky"],
        ["미국 남성 (am_)", "am_adam, am_echo, am_eric, am_fenrir, am_liam, am_michael, am_onyx, am_puck, am_santa"],
        ["영국 여성 (bf_)", "bf_alice, bf_emma, bf_isabella, bf_lily"],
        ["영국 남성 (bm_)", "bm_daniel, bm_fable, bm_george, bm_lewis"],
    ], [3.6, 13])
    d.p("`[speaker: student]` 문단은 기본으로 `am_michael`(속도 1.05)로 읽습니다.", size=9.5, color=GRAY)
    d.save(out)


def main():
    global REPO
    if len(sys.argv) > 1:
        REPO = sys.argv[1].rstrip("/")
    shot = HERE / "images" / "main_window.png"
    deploy_manual(HERE / "LectureTTS_배포매뉴얼.docx")
    user_manual(HERE / "LectureTTS_사용설명서.docx", shot)


if __name__ == "__main__":
    main()
